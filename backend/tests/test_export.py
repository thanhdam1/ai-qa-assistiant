import io
import openpyxl

from app import models
from app.services import export


def test_csv_exports_are_utf8_bom_and_preserve_multiline(review_content, test_content):
    review = models.RequirementReview(content=review_content.model_dump(mode="json"))
    test_set = models.TestCaseSet(content=test_content.model_dump(mode="json"))
    review_bytes = export.review_csv(review)
    tests_bytes = export.test_set_csv(test_set)
    assert review_bytes.startswith(b"\xef\xbb\xbf")
    assert "Người dùng đăng nhập" in review_bytes.decode("utf-8-sig")
    assert tests_bytes.startswith(b"\xef\xbb\xbf")
    assert "TC-001" in tests_bytes.decode("utf-8-sig")


def test_excel_export_test_set_generates_valid_workbook(test_content):
    # Add mock coverage matrix to test_content
    test_dict = test_content.model_dump(mode="json")
    test_dict["coverage_matrix"] = [
        {
            "target_type": "acceptance_criterion",
            "target_id": "AC-001",
            "status": "covered",
            "test_case_ids": ["TC-001", "TC-002"],
            "explanation": "Đã phủ đầy đủ điều kiện đăng nhập",
            "assertion_evidence": [
                {
                    "test_case_id": "TC-001",
                    "assertion": "Xác nhận mã 200 trả về",
                    "verifies_business_outcome": True,
                }
            ],
        }
    ]

    test_set = models.TestCaseSet(
        id="test-set-1234",
        story_id="story-5678",
        story_version=1,
        status="approved",
        content=test_dict,
        case_readiness=[{"case_id": "TC-001", "readiness": "ready", "notes": "Approved"}],
    )

    xlsx_bytes = export.test_set_excel(test_set)
    assert xlsx_bytes is not None
    assert len(xlsx_bytes) > 2000

    # Parse back with openpyxl to verify validity
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet_names = wb.sheetnames
    assert "Danh sách Test Case" in sheet_names
    assert "Ma trận Phủ (Coverage)" in sheet_names
    assert "Tổng quan & KPI" in sheet_names

    ws_cases = wb["Danh sách Test Case"]
    assert ws_cases["A1"].value == "BÁO CÁO ĐẶC TẢ TEST CASE (TEST SPECIFICATION REPORT)"
    # Row 5 is header, row 6 is first test case
    assert ws_cases.cell(row=5, column=2).value == "Mã Test Case"
    assert ws_cases.cell(row=6, column=2).value == "TC-001"
    assert ws_cases.cell(row=6, column=18).value == "ready"  # Readiness status

    ws_cov = wb["Ma trận Phủ (Coverage)"]
    assert ws_cov.cell(row=4, column=3).value == "AC-001"
    assert ws_cov.cell(row=4, column=4).value == "COVERED"

    ws_sum = wb["Tổng quan & KPI"]
    assert "THÔNG TIN CHUNG BỘ KIỂM THỬ" in str(ws_sum["A3"].value)


def test_excel_export_zero_gap_suite_generates_valid_workbook():
    mock_suite = {
        "effective_coverage_pct": 100.0,
        "verification_mode": "3-Way Zero-Gap (Story + OpenAPI + AST CFG)",
        "retries_used": 1,
        "test_cases": [
            {
                "test_case_id": "TC-REQ-001",
                "title": "Kiểm tra tạo giao dịch hợp lệ",
                "category": "Functional",
                "priority": "High",
                "resolution_status": "COVERED_AUTOMATED",
                "preconditions": ["User balance > 1000"],
                "test_data_payload": {"amount": 500, "currency": "VND"},
                "test_steps": ["Gửi POST /transactions", "Kiểm tra số dư mới"],
                "expected_result": "Giao dịch thành công 200 OK",
                "expected_http_status": 200,
                "covered_req_ids": ["REQ-01"],
                "covered_bva_ids": ["BVA-AMOUNT-NORMAL"],
                "covered_ast_ids": ["CFG-BRANCH-12"],
            }
        ],
        "gaps_for_qa": [
            {
                "target_id": "BVA-AMOUNT-MAX",
                "gap_type": "SPEC_CONFLICT",
                "reason": "Story không giới hạn số tiền tối đa nhưng OpenAPI giới hạn 50,000,000 VND",
                "status": "NEEDS_CLARIFICATION",
            }
        ],
    }

    xlsx_bytes = export.zero_gap_suite_excel(mock_suite, story_title="Chuyển tiền nhanh 24/7")
    assert xlsx_bytes is not None
    assert len(xlsx_bytes) > 2000

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Zero-Gap Test Cases" in wb.sheetnames
    assert "Khoảng trống QA (Gaps)" in wb.sheetnames

    ws_cases = wb["Zero-Gap Test Cases"]
    assert "BÁO CÁO KIỂM THỬ ZERO-GAP" in str(ws_cases["A1"].value)
    assert ws_cases.cell(row=6, column=2).value == "TC-REQ-001"
    assert ws_cases.cell(row=6, column=11).value == "200"

    ws_gaps = wb["Khoảng trống QA (Gaps)"]
    assert ws_gaps.cell(row=4, column=2).value == "BVA-AMOUNT-MAX"
    assert ws_gaps.cell(row=4, column=3).value == "SPEC_CONFLICT"
