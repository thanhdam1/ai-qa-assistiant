import sys
import os
import json
import pytest
import openpyxl
import io

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.parsers.prd_parser import PRDParser
from app.agents.test_case_agent import TestCaseGenerator
from app.agents.rtm_agent import RTMAgent
from app.agents.synthetic_data import SyntheticDataGenerator
from app.agents.export_service import ExportService
from app.api.v1.qa_assistant import generate_test_cases, generate_synthetic_data, export_excel, export_csv, GenerateTestRequest, SyntheticDataRequest

SAMPLE_PRD = """
# User Authentication PRD

## User Story: Login
As a customer, I want to sign in with email and password, so that I can access my account.

### Acceptance Criteria:
- Given valid registered email and password When user clicks login Then redirect to dashboard.
- Given invalid password When user clicks login Then display error "Invalid credentials".
- Given email field is empty When user clicks login Then display validation error "Email is required".
"""

def test_synthetic_data_generator():
    schema = {
        "type": "object",
        "required": ["email", "age"],
        "properties": {
            "email": {"type": "string", "format": "email"},
            "age": {"type": "integer", "minimum": 18, "maximum": 99},
            "role": {"type": "string", "enum": ["admin", "user"]}
        }
    }
    dataset = SyntheticDataGenerator.generate_for_schema(schema, "/api/auth/register", "POST")
    
    assert dataset.endpoint_path == "/api/auth/register"
    assert len(dataset.valid_payloads) >= 1
    # Check valid payload properties
    valid_p = dataset.valid_payloads[0]
    assert "@" in valid_p["email"]
    assert 18 <= valid_p["age"] <= 99
    assert valid_p["role"] in ["admin", "user"]

    # Check boundary
    assert len(dataset.boundary_payloads) >= 1
    # Check security
    assert len(dataset.security_payloads) >= 1
    assert any("' OR" in str(v) or "<script>" in str(v) for v in dataset.security_payloads[0].values())

def test_test_case_generation_and_rtm():
    prd = PRDParser.parse_markdown(SAMPLE_PRD, "Auth_PRD.md")
    test_suite = TestCaseGenerator.generate_from_prd(prd)

    assert test_suite.total_test_cases >= 4
    categories = [tc.category for tc in test_suite.test_cases]
    assert "Positive" in categories
    assert "Boundary Value Analysis" in categories
    assert "Negative" in categories
    assert "Security" in categories

    # Test RTM Matrix
    rtm = RTMAgent.build_matrix(prd, test_suite)
    assert rtm.total_requirements == 3
    assert rtm.covered_requirements == 3
    assert rtm.coverage_percentage == 100.0
    assert len(rtm.uncovered_gaps) == 0

def test_export_service():
    prd = PRDParser.parse_markdown(SAMPLE_PRD, "Auth_PRD.md")
    test_suite = TestCaseGenerator.generate_from_prd(prd)

    # Test CSV
    csv_str = ExportService.to_csv(test_suite)
    assert "Test ID" in csv_str
    assert "Positive" in csv_str

    # Test Excel
    excel_bytes = ExportService.to_excel_bytes(test_suite)
    assert len(excel_bytes) > 1000
    # Verify openpyxl can read the generated workbook
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb["Test Cases"]
    assert ws.cell(row=1, column=1).value == "Test ID"
    assert ws.max_row >= 5

def test_qa_api_endpoints():
    req = GenerateTestRequest(prd_content=SAMPLE_PRD, filename="Auth_PRD.md")
    res = generate_test_cases(req)
    
    assert res.test_suite.total_test_cases >= 4
    assert res.rtm_matrix.coverage_percentage == 100.0

    # Test synthetic data API
    synth_req = SyntheticDataRequest(
        schema_definition={
            "type": "object",
            "properties": {"name": {"type": "string"}, "score": {"type": "number", "minimum": 0}}
        }
    )
    synth_res = generate_synthetic_data(synth_req)
    assert len(synth_res.valid_payloads) >= 1

def test_real_live_runner_and_websocket():
    from fastapi.testclient import TestClient
    from app.main import app

    client = TestClient(app)

    # 1. Real code with valid syntax and assertions
    valid_code = """
def test_addition():
    assert 2 + 2 == 4

def test_string():
    assert "hello".upper() == "HELLO"
"""
    with client.websocket_connect("/api/v1/qa/ws/runner") as ws:
        ws.send_text(json.dumps({"code": valid_code, "framework": "playwright"}))
        received = []
        while True:
            try:
                msg = ws.receive_text()
                received.append(msg)
                if "[STATUS] PASSED" in msg or "[STATUS] FAILED" in msg:
                    break
            except Exception:
                break
        
        all_text = "\n".join(received)
        assert "[STATUS] PASSED" in all_text
        assert "Phân tích cây cú pháp trừu tượng AST" in all_text
        assert "Biên dịch bytecode với py_compile" in all_text

    # 2. Broken code with AST Syntax error
    broken_syntax_code = """
def broken_func()
    print("missing colon")
"""
    with client.websocket_connect("/api/v1/qa/ws/runner") as ws:
        ws.send_text(json.dumps({"code": broken_syntax_code, "framework": "playwright"}))
        received = []
        while True:
            try:
                msg = ws.receive_text()
                received.append(msg)
                if "[STATUS] PASSED" in msg or "[STATUS] FAILED" in msg:
                    break
            except Exception:
                break
        
        all_text = "\n".join(received)
        assert "[STATUS] FAILED" in all_text
        assert "Phát hiện lỗi cú pháp AST" in all_text

