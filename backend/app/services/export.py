import csv
import datetime
import io
import json
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import ApiContractReview, RequirementReview, TestCaseSet


def json_bytes(content: dict) -> bytes:
    return json.dumps(content, ensure_ascii=False, indent=2).encode("utf-8")


def review_csv(review: RequirementReview) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["section", "id", "title", "detail", "severity", "source_reference", "classification"])
    content = review.content
    writer.writerow(["business_summary", "", "", content["business_summary"], "", "", "requirement"])
    for item in content.get("existing_acceptance_criteria", []):
        writer.writerow(
            [
                "existing_acceptance_criteria",
                item["id"],
                "",
                item["criterion"],
                "",
                item["source_reference"],
                "requirement",
            ]
        )
    for section in (
        "missing_acceptance_criteria",
        "ambiguities_and_conflicts",
        "stakeholder_questions",
        "edge_cases",
        "assumptions",
        "risks",
    ):
        for item in content.get(section, []):
            writer.writerow(
                [
                    section,
                    item["id"],
                    item["title"],
                    item["detail"],
                    item["severity"],
                    item["source_reference"],
                    item.get("classification", "recommendation"),
                ]
            )
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def api_contract_review_csv(review: ApiContractReview) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(
        [
            "id",
            "operation_id",
            "method",
            "path",
            "type",
            "title",
            "detail",
            "severity",
            "classification",
            "source_reference",
            "story_reference",
            "expected_status_codes",
            "needs_clarification",
        ]
    )
    content = review.content
    for finding in content.get("findings", []):
        writer.writerow(
            [
                finding["id"],
                finding["operation_id"],
                finding["method"],
                finding["path"],
                finding["type"],
                finding["title"],
                finding["detail"],
                finding["severity"],
                finding["classification"],
                finding["source_reference"],
                finding.get("story_reference") or "",
                ",".join(str(code) for code in finding.get("expected_status_codes", [])),
                str(finding["needs_clarification"]).lower(),
            ]
        )
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def test_set_csv(test_set: TestCaseSet) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(
        [
            "id",
            "title",
            "objective",
            "type",
            "basis",
            "priority",
            "preconditions",
            "test_data",
            "steps",
            "expected_result",
            "requirement_traceability",
            "assumptions",
            "needs_clarification",
            "api_method",
            "api_path",
            "api_request",
            "expected_status_codes",
            "api_traceability",
            "coverage_target_type",
            "coverage_target_id",
            "coverage_status",
            "coverage_test_case_ids",
            "coverage_explanation",
        ]
    )
    for case in test_set.content.get("test_cases", []):
        steps = "\n".join(
            f"{step['order']}. {step['action']} => {step['expected_result']}"
            for step in case["steps"]
        )
        writer.writerow(
            [
                case["id"],
                case["title"],
                case["objective"],
                case["type"],
                case.get("basis", "recommendation"),
                case["priority"],
                "\n".join(case["preconditions"]),
                "\n".join(case["test_data"]),
                steps,
                case["expected_result"],
                "\n".join(case["requirement_traceability"]),
                "\n".join(case["assumptions"]),
                str(case["needs_clarification"]).lower(),
                case.get("api_method") or "",
                case.get("api_path") or "",
                case.get("api_request") or "",
                ",".join(str(code) for code in case.get("expected_status_codes", [])),
                "\n".join(case.get("api_traceability", [])),
                "",
                "",
                "",
                "",
                "",
            ]
        )
    for entry in test_set.content.get("coverage_matrix", []):
        writer.writerow(
            ["coverage"]
            + [""] * 17
            + [
                entry["target_type"],
                entry["target_id"],
                entry["status"],
                "\n".join(entry.get("test_case_ids", [])),
                entry["explanation"],
            ]
        )
    return ("\ufeff" + output.getvalue()).encode("utf-8")


# ============================================================================
# EXCEL EXPORT ENGINE - ISO/IEC/IEEE 29119-3 & ISTQB FORMATTING
# ============================================================================

FONT_NAME = "Segoe UI"
COLOR_PRIMARY = "1E3A8A"        # Deep Corporate Navy
COLOR_SECONDARY = "0F766E"      # Deep Emerald Teal
COLOR_BORDER = "CBD5E1"         # Slate 300
COLOR_ZEBRA = "F8FAFC"          # Slate 50
COLOR_HEADER_BG = "1E3A8A"
COLOR_INFO_BG = "F1F5F9"

THIN_BORDER = Border(
    left=Side(style="thin", color=COLOR_BORDER),
    right=Side(style="thin", color=COLOR_BORDER),
    top=Side(style="thin", color=COLOR_BORDER),
    bottom=Side(style="thin", color=COLOR_BORDER),
)

PRIORITY_COLOR_MAP = {
    "CRITICAL": {
        "fill": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "font": Font(name=FONT_NAME, size=9.5, bold=True, color="991B1B"),
    },
    "HIGH": {
        "fill": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "font": Font(name=FONT_NAME, size=9.5, bold=True, color="B91C1C"),
    },
    "MEDIUM": {
        "fill": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "font": Font(name=FONT_NAME, size=9.5, bold=True, color="B45309"),
    },
    "LOW": {
        "fill": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "font": Font(name=FONT_NAME, size=9.5, bold=True, color="15803D"),
    },
}

STATUS_COLOR_MAP = {
    "COVERED": {
        "fill": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "font": Font(name=FONT_NAME, size=9.5, bold=True, color="15803D"),
    },
    "NEEDS_CLARIFICATION": {
        "fill": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "font": Font(name=FONT_NAME, size=9.5, bold=True, color="B45309"),
    },
    "UNCOVERED": {
        "fill": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "font": Font(name=FONT_NAME, size=9.5, bold=True, color="B91C1C"),
    },
    "COVERED_AUTOMATED": {
        "fill": PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),
        "font": Font(name=FONT_NAME, size=9.5, bold=True, color="1E40AF"),
    },
}


def _format_bullets(items: Any) -> str:
    if not items:
        return ""
    if isinstance(items, str):
        return items
    if isinstance(items, list):
        return "\n".join(f"• {str(item).strip()}" for item in items if str(item).strip())
    return str(items)


def _format_steps(steps: Any) -> str:
    if not steps:
        return ""
    if isinstance(steps, str):
        return steps
    formatted = []
    for idx, step in enumerate(steps, 1):
        if isinstance(step, dict):
            order = step.get("order", idx)
            action = step.get("action", "")
            exp = step.get("expected_result", "")
            if exp:
                formatted.append(f"{order}. {action}\n   -> Kỳ vọng: {exp}")
            else:
                formatted.append(f"{order}. {action}")
        elif isinstance(step, str):
            formatted.append(step if step.startswith(f"{idx}.") else f"{idx}. {step}")
        else:
            action = getattr(step, "action", str(step))
            exp = getattr(step, "expected_result", "")
            order = getattr(step, "order", idx)
            if exp:
                formatted.append(f"{order}. {action}\n   -> Kỳ vọng: {exp}")
            else:
                formatted.append(f"{order}. {action}")
    return "\n".join(formatted)


def _format_test_data(data: Any) -> str:
    if not data:
        return ""
    if isinstance(data, dict):
        return json.dumps(data, ensure_ascii=False, indent=2)
    if isinstance(data, list):
        if all(isinstance(x, str) for x in data):
            return "\n".join(f"• {x}" for x in data)
        return json.dumps(data, ensure_ascii=False, indent=2)
    return str(data)


def _auto_fit_columns(ws, max_cols: int, custom_widths: dict[int, int] | None = None):
    custom_widths = custom_widths or {}
    for col_idx in range(1, max_cols + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx in custom_widths:
            ws.column_dimensions[col_letter].width = custom_widths[col_idx]
            continue
        max_len = 0
        for cell in ws[col_letter]:
            if cell.row in (1, 2, 3, 4):  # Skip banner / metadata rows
                continue
            val = str(cell.value or "")
            lines = val.split("\n")
            line_len = max(len(line) for line in lines) if lines else 0
            if line_len > max_len:
                max_len = line_len
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 12)


def test_set_excel(test_set: TestCaseSet) -> bytes:
    wb = openpyxl.Workbook()
    now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    # Map readiness assessments if present
    readiness_map = {}
    for r in getattr(test_set, "case_readiness", []) or []:
        if isinstance(r, dict):
            readiness_map[r.get("case_id")] = r.get("readiness", "not_assessed")

    story_title = ""
    if getattr(test_set, "story", None) and getattr(test_set.story, "title", None):
        story_title = test_set.story.title
    else:
        story_title = f"Story #{test_set.story_id}"

    content = test_set.content or {}
    test_cases = content.get("test_cases", [])
    coverage_matrix = content.get("coverage_matrix", [])

    # =========================================================================
    # SHEET 1: DANH SÁCH TEST CASE
    # =========================================================================
    ws_cases = wb.active
    ws_cases.title = "Danh sách Test Case"
    ws_cases.views.sheetView[0].showGridLines = True

    # Row 1: Banner Header
    ws_cases.merge_cells("A1:U1")
    top_cell = ws_cases["A1"]
    top_cell.value = "BÁO CÁO ĐẶC TẢ TEST CASE (TEST SPECIFICATION REPORT)"
    top_cell.font = Font(name=FONT_NAME, size=15, bold=True, color="FFFFFF")
    top_cell.fill = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    top_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_cases.row_dimensions[1].height = 38

    # Row 2 & 3: Metadata info block
    metadata_fields = [
        ("A2", "Dự án / Hệ thống:", "B2", "AI QA Assistant"),
        ("E2", "User Story:", "F2", story_title),
        ("K2", "Ngày xuất báo cáo:", "L2", now_str),
        ("A3", "Tổng số Test Case:", "B3", len(test_cases)),
        ("E3", "Trạng thái bộ test:", "F3", (test_set.status or "draft").upper()),
        ("K3", "Độ tin cậy tổng thể:", "L3", str(content.get("overall_confidence_score", "N/A"))),
    ]
    for lbl_pos, lbl_val, val_pos, val_val in metadata_fields:
        c_lbl = ws_cases[lbl_pos]
        c_lbl.value = lbl_val
        c_lbl.font = Font(name=FONT_NAME, size=9.5, bold=True, color="334155")
        c_val = ws_cases[val_pos]
        c_val.value = val_val
        c_val.font = Font(name=FONT_NAME, size=9.5, color="0F172A")

    ws_cases.row_dimensions[2].height = 20
    ws_cases.row_dimensions[3].height = 20
    ws_cases.row_dimensions[4].height = 8  # Spacer row

    # Row 5: Column Headers
    headers_cases = [
        "STT",
        "Mã Test Case",
        "Tiêu đề Test Case",
        "Mục tiêu kiểm thử",
        "Phân loại",
        "Độ ưu tiên",
        "Cơ sở",
        "Điều kiện tiên quyết",
        "Dữ liệu kiểm thử",
        "Các bước thực hiện",
        "Kết quả mong đợi",
        "API Method",
        "API Path",
        "Mã HTTP mong đợi",
        "Truy vết Yêu cầu",
        "Truy vết API",
        "Cần làm rõ?",
        "Đánh giá sẵn sàng",
        "Kết quả thực tế",
        "Trạng thái thực thi",
        "Ghi chú",
    ]

    ws_cases.row_dimensions[5].height = 28
    for col_idx, header_text in enumerate(headers_cases, 1):
        cell = ws_cases.cell(row=5, column=col_idx, value=header_text)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    # Data Rows
    current_row = 6
    for idx, tc in enumerate(test_cases, 1):
        is_even = (idx % 2 == 0)
        row_fill = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid") if is_even else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        prio_raw = str(tc.get("priority", "medium")).upper()
        prio_style = PRIORITY_COLOR_MAP.get(prio_raw, PRIORITY_COLOR_MAP["MEDIUM"])

        basis_text = "Yêu cầu (SRS)" if tc.get("basis") == "requirement" else "Gợi ý bổ sung"
        clarif_text = "Có" if tc.get("needs_clarification") else "Không"
        readiness_val = readiness_map.get(tc.get("id"), "Chưa đánh giá")

        http_codes = ",".join(str(code) for code in tc.get("expected_status_codes", []))
        req_trace = "\n".join(tc.get("requirement_traceability", []))
        api_trace = "\n".join(tc.get("api_traceability", []))

        row_values = [
            (idx, Alignment(horizontal="center", vertical="center"), None, None),
            (tc.get("id", ""), Alignment(horizontal="center", vertical="center"), Font(name=FONT_NAME, size=9.5, bold=True), None),
            (tc.get("title", ""), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (tc.get("objective", ""), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (str(tc.get("type", "")).capitalize(), Alignment(horizontal="center", vertical="center"), None, None),
            (prio_raw, Alignment(horizontal="center", vertical="center"), prio_style["font"], prio_style["fill"]),
            (basis_text, Alignment(horizontal="center", vertical="center"), None, None),
            (_format_bullets(tc.get("preconditions")), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (_format_test_data(tc.get("test_data")), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (_format_steps(tc.get("steps")), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (tc.get("expected_result", ""), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (tc.get("api_method") or "", Alignment(horizontal="center", vertical="center"), None, None),
            (tc.get("api_path") or "", Alignment(horizontal="left", vertical="center"), None, None),
            (http_codes, Alignment(horizontal="center", vertical="center"), None, None),
            (req_trace, Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (api_trace, Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (clarif_text, Alignment(horizontal="center", vertical="center"), None, None),
            (readiness_val, Alignment(horizontal="center", vertical="center"), None, None),
            ("", Alignment(horizontal="left", vertical="top"), None, None),              # Actual Result
            ("Chưa chạy", Alignment(horizontal="center", vertical="center"), None, None), # Execution status
            ("", Alignment(horizontal="left", vertical="top"), None, None),              # Notes
        ]

        for col_idx, (val, align, font_override, fill_override) in enumerate(row_values, 1):
            cell = ws_cases.cell(row=current_row, column=col_idx, value=val)
            cell.alignment = align
            cell.font = font_override or Font(name=FONT_NAME, size=9.5, color="0F172A")
            cell.fill = fill_override or row_fill
            cell.border = THIN_BORDER

        current_row += 1

    # Freeze panes and Auto Filter
    ws_cases.freeze_panes = "A6"
    if current_row > 6:
        ws_cases.auto_filter.ref = f"A5:U{current_row - 1}"

    # Custom column widths for readability
    cases_col_widths = {
        1: 6,    # STT
        2: 15,   # Mã TC
        3: 32,   # Tiêu đề
        4: 28,   # Mục tiêu
        5: 14,   # Phân loại
        6: 13,   # Độ ưu tiên
        7: 15,   # Cơ sở
        8: 28,   # Tiền đề
        9: 25,   # Test data
        10: 42,  # Steps
        11: 34,  # Expected result
        12: 12,  # Method
        13: 24,  # Path
        14: 12,  # HTTP code
        15: 18,  # Req trace
        16: 18,  # API trace
        17: 12,  # Clarification
        18: 16,  # Readiness
        19: 24,  # Actual result
        20: 16,  # Exec status
        21: 20,  # Notes
    }
    _auto_fit_columns(ws_cases, len(headers_cases), cases_col_widths)

    # =========================================================================
    # SHEET 2: MA TRẬN ĐỘ PHỦ YÊU CẦU & RỦI RO (TRACEABILITY MATRIX)
    # =========================================================================
    if coverage_matrix:
        ws_cov = wb.create_sheet(title="Ma trận Phủ (Coverage)")
        ws_cov.views.sheetView[0].showGridLines = True

        # Banner
        ws_cov.merge_cells("A1:G1")
        top_cov = ws_cov["A1"]
        top_cov.value = "MA TRẬN ĐỘ PHỦ YÊU CẦU & BẰNG CHỨNG XÁC NHẬN (TRACEABILITY MATRIX)"
        top_cov.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
        top_cov.fill = PatternFill(start_color=COLOR_SECONDARY, end_color=COLOR_SECONDARY, fill_type="solid")
        top_cov.alignment = Alignment(horizontal="center", vertical="center")
        ws_cov.row_dimensions[1].height = 36

        headers_cov = [
            "STT",
            "Loại mục tiêu (Target Type)",
            "Mã mục tiêu (Target ID)",
            "Trạng thái phủ (Coverage Status)",
            "Test Case liên kết (Linked Test Cases)",
            "Giải trình độ phủ (Explanation)",
            "Bằng chứng xác nhận (Assertion Evidence)",
        ]

        ws_cov.row_dimensions[3].height = 26
        for col_idx, header_text in enumerate(headers_cov, 1):
            cell = ws_cov.cell(row=3, column=col_idx, value=header_text)
            cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=COLOR_SECONDARY, end_color=COLOR_SECONDARY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN_BORDER

        cov_row = 4
        for idx, entry in enumerate(coverage_matrix, 1):
            is_even = (idx % 2 == 0)
            row_fill = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid") if is_even else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

            status_raw = str(entry.get("status", "covered")).upper()
            status_style = STATUS_COLOR_MAP.get(status_raw, STATUS_COLOR_MAP["COVERED"])

            evidence_items = entry.get("assertion_evidence", [])
            evidence_str = "\n".join(
                f"• [{ev.get('test_case_id')}]: {ev.get('assertion')}"
                for ev in evidence_items
                if isinstance(ev, dict)
            )

            tc_ids = ", ".join(entry.get("test_case_ids", []))

            cov_values = [
                (idx, Alignment(horizontal="center", vertical="center"), None, None),
                (entry.get("target_type", ""), Alignment(horizontal="center", vertical="center"), None, None),
                (entry.get("target_id", ""), Alignment(horizontal="center", vertical="center"), Font(name=FONT_NAME, size=9.5, bold=True), None),
                (status_raw, Alignment(horizontal="center", vertical="center"), status_style["font"], status_style["fill"]),
                (tc_ids, Alignment(horizontal="left", vertical="center", wrap_text=True), None, None),
                (entry.get("explanation", ""), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
                (evidence_str, Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            ]

            for col_idx, (val, align, font_override, fill_override) in enumerate(cov_values, 1):
                cell = ws_cov.cell(row=cov_row, column=col_idx, value=val)
                cell.alignment = align
                cell.font = font_override or Font(name=FONT_NAME, size=9.5, color="0F172A")
                cell.fill = fill_override or row_fill
                cell.border = THIN_BORDER

            cov_row += 1

        ws_cov.freeze_panes = "A4"
        if cov_row > 4:
            ws_cov.auto_filter.ref = f"A3:G{cov_row - 1}"

        cov_widths = {1: 6, 2: 24, 3: 16, 4: 18, 5: 26, 6: 42, 7: 45}
        _auto_fit_columns(ws_cov, len(headers_cov), cov_widths)

    # =========================================================================
    # SHEET 3: DASHBOARD & THỐNG KÊ (SUMMARY)
    # =========================================================================
    ws_sum = wb.create_sheet(title="Tổng quan & KPI")
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum.merge_cells("A1:F1")
    top_sum = ws_sum["A1"]
    top_sum.value = "TỔNG QUAN CHẤT LƯỢNG & PHÂN BỔ BỘ TEST (TEST SUITE DASHBOARD)"
    top_sum.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
    top_sum.fill = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    top_sum.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 36

    # Summary Statistics
    total_cases = len(test_cases)
    prio_counts = {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 0, "LOW": 0}
    type_counts: dict[str, int] = {}
    for tc in test_cases:
        p = str(tc.get("priority", "medium")).upper()
        prio_counts[p] = prio_counts.get(p, 0) + 1
        t = str(tc.get("type", "functional")).capitalize()
        type_counts[t] = type_counts.get(t, 0) + 1

    # Table 1: Thông tin bộ kiểm thử
    ws_sum["A3"] = "THÔNG TIN CHUNG BỘ KIỂM THỬ"
    ws_sum["A3"].font = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_PRIMARY)

    info_rows = [
        ("Tên User Story / Tính năng:", story_title),
        ("Mã Story ID:", test_set.story_id),
        ("Phiên bản Story:", str(test_set.story_version)),
        ("Mã Bộ Test (Test Set ID):", test_set.id),
        ("Trạng thái phê duyệt:", (test_set.status or "draft").upper()),
        ("Độ tin cậy AI (Confidence):", str(content.get("overall_confidence_score", "N/A"))),
        ("Thời gian xuất báo cáo:", now_str),
        ("Khuyến nghị thực thi:", str(content.get("execution_recommendation", "N/A"))),
    ]

    for idx, (label, val) in enumerate(info_rows, 4):
        c_lbl = ws_sum.cell(row=idx, column=1, value=label)
        c_lbl.font = Font(name=FONT_NAME, size=9.5, bold=True, color="334155")
        c_lbl.fill = PatternFill(start_color=COLOR_INFO_BG, end_color=COLOR_INFO_BG, fill_type="solid")
        c_lbl.border = THIN_BORDER

        ws_sum.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=3)
        c_val = ws_sum.cell(row=idx, column=2, value=val)
        c_val.font = Font(name=FONT_NAME, size=9.5, color="0F172A")
        c_val.border = THIN_BORDER
        ws_sum.cell(row=idx, column=3).border = THIN_BORDER

    # Table 2: Phân bổ theo Độ ưu tiên
    prio_start_row = 4
    ws_sum.cell(row=prio_start_row - 1, column=5, value="PHÂN BỔ THEO ĐỘ ƯU TIÊN").font = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_PRIMARY)
    ws_sum.cell(row=prio_start_row, column=5, value="Độ ưu tiên").font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    ws_sum.cell(row=prio_start_row, column=5).fill = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    ws_sum.cell(row=prio_start_row, column=5).border = THIN_BORDER

    ws_sum.cell(row=prio_start_row, column=6, value="Số lượng").font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    ws_sum.cell(row=prio_start_row, column=6).fill = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    ws_sum.cell(row=prio_start_row, column=6).border = THIN_BORDER

    r_offset = 1
    for p_name, count in prio_counts.items():
        if count > 0 or p_name in ("HIGH", "MEDIUM", "LOW"):
            cell_p = ws_sum.cell(row=prio_start_row + r_offset, column=5, value=p_name)
            p_style = PRIORITY_COLOR_MAP.get(p_name, PRIORITY_COLOR_MAP["MEDIUM"])
            cell_p.font = p_style["font"]
            cell_p.fill = p_style["fill"]
            cell_p.alignment = Alignment(horizontal="center", vertical="center")
            cell_p.border = THIN_BORDER

            cell_c = ws_sum.cell(row=prio_start_row + r_offset, column=6, value=count)
            cell_c.font = Font(name=FONT_NAME, size=9.5, bold=True, color="0F172A")
            cell_c.alignment = Alignment(horizontal="center", vertical="center")
            cell_c.border = THIN_BORDER
            r_offset += 1

    # Total row
    cell_tot_lbl = ws_sum.cell(row=prio_start_row + r_offset, column=5, value="TỔNG CỘNG")
    cell_tot_lbl.font = Font(name=FONT_NAME, size=9.5, bold=True, color="0F172A")
    cell_tot_lbl.alignment = Alignment(horizontal="center", vertical="center")
    cell_tot_lbl.border = THIN_BORDER

    cell_tot_val = ws_sum.cell(row=prio_start_row + r_offset, column=6, value=total_cases)
    cell_tot_val.font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_PRIMARY)
    cell_tot_val.alignment = Alignment(horizontal="center", vertical="center")
    cell_tot_val.border = THIN_BORDER

    ws_sum.column_dimensions["A"].width = 28
    ws_sum.column_dimensions["B"].width = 30
    ws_sum.column_dimensions["C"].width = 20
    ws_sum.column_dimensions["D"].width = 6
    ws_sum.column_dimensions["E"].width = 18
    ws_sum.column_dimensions["F"].width = 14

    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def zero_gap_suite_excel(suite: Any, story_title: str = "Zero-Gap Test Suite") -> bytes:
    wb = openpyxl.Workbook()
    now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    # Handle Pydantic model or dict
    if hasattr(suite, "model_dump"):
        suite_dict = suite.model_dump()
    elif isinstance(suite, dict):
        suite_dict = suite
    else:
        suite_dict = {}

    test_cases = suite_dict.get("test_cases", [])
    gaps_for_qa = suite_dict.get("gaps_for_qa", [])
    cov_pct = suite_dict.get("effective_coverage_pct", 100.0)
    verif_mode = suite_dict.get("verification_mode", "3-Way Zero-Gap (Story + OpenAPI + AST CFG)")
    retries = suite_dict.get("retries_used", 0)

    # =========================================================================
    # SHEET 1: ZERO-GAP TEST CASES
    # =========================================================================
    ws_cases = wb.active
    ws_cases.title = "Zero-Gap Test Cases"
    ws_cases.views.sheetView[0].showGridLines = True

    # Banner Header
    ws_cases.merge_cells("A1:Q1")
    top_cell = ws_cases["A1"]
    top_cell.value = "BÁO CÁO KIỂM THỬ ZERO-GAP v2.0 (100% MATHEMATICAL COVERAGE)"
    top_cell.font = Font(name=FONT_NAME, size=15, bold=True, color="FFFFFF")
    top_cell.fill = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    top_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_cases.row_dimensions[1].height = 38

    # Metadata rows
    metadata_fields = [
        ("A2", "User Story / Module:", "B2", story_title),
        ("F2", "Chế độ xác minh:", "G2", verif_mode),
        ("K2", "Ngày xuất:", "L2", now_str),
        ("A3", "Độ phủ hiệu dụng:", "B3", f"{cov_pct:.1f}%"),
        ("F3", "Số lần Self-Healing:", "G3", str(retries)),
        ("K3", "Tổng số Test Case:", "L3", len(test_cases)),
    ]
    for lbl_pos, lbl_val, val_pos, val_val in metadata_fields:
        c_lbl = ws_cases[lbl_pos]
        c_lbl.value = lbl_val
        c_lbl.font = Font(name=FONT_NAME, size=9.5, bold=True, color="334155")
        c_val = ws_cases[val_pos]
        c_val.value = val_val
        c_val.font = Font(name=FONT_NAME, size=9.5, color="0F172A")

    ws_cases.row_dimensions[2].height = 20
    ws_cases.row_dimensions[3].height = 20
    ws_cases.row_dimensions[4].height = 8

    # Column Headers
    headers_zero = [
        "STT",
        "Mã Test Case",
        "Tiêu đề Test Case",
        "Phân loại",
        "Độ ưu tiên",
        "Trạng thái xử lý",
        "Điều kiện tiên quyết",
        "Dữ liệu Payload",
        "Các bước thực hiện",
        "Kết quả mong đợi",
        "Mã HTTP",
        "Truy vết Yêu cầu SRS",
        "Truy vết Biên OpenAPI BVA",
        "Truy vết Nhánh Code AST",
        "Kết quả thực tế",
        "Trạng thái thực thi",
        "Ghi chú",
    ]

    ws_cases.row_dimensions[5].height = 28
    for col_idx, header_text in enumerate(headers_zero, 1):
        cell = ws_cases.cell(row=5, column=col_idx, value=header_text)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    current_row = 6
    for idx, tc in enumerate(test_cases, 1):
        is_even = (idx % 2 == 0)
        row_fill = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid") if is_even else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        prio_raw = str(tc.get("priority", "HIGH")).upper()
        prio_style = PRIORITY_COLOR_MAP.get(prio_raw, PRIORITY_COLOR_MAP["HIGH"])

        resol_raw = str(tc.get("resolution_status", "COVERED_AUTOMATED")).upper()
        resol_style = STATUS_COLOR_MAP.get(resol_raw, STATUS_COLOR_MAP["COVERED_AUTOMATED"])

        req_ids = "\n".join(tc.get("covered_req_ids", []))
        bva_ids = "\n".join(tc.get("covered_bva_ids", []))
        ast_ids = "\n".join(tc.get("covered_ast_ids", []))

        row_values = [
            (idx, Alignment(horizontal="center", vertical="center"), None, None),
            (tc.get("test_case_id", ""), Alignment(horizontal="center", vertical="center"), Font(name=FONT_NAME, size=9.5, bold=True), None),
            (tc.get("title", ""), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (str(tc.get("category", "")).capitalize(), Alignment(horizontal="center", vertical="center"), None, None),
            (prio_raw, Alignment(horizontal="center", vertical="center"), prio_style["font"], prio_style["fill"]),
            (resol_raw, Alignment(horizontal="center", vertical="center"), resol_style["font"], resol_style["fill"]),
            (_format_bullets(tc.get("preconditions")), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (_format_test_data(tc.get("test_data_payload")), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (_format_steps(tc.get("test_steps")), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (tc.get("expected_result", ""), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (str(tc.get("expected_http_status", "")), Alignment(horizontal="center", vertical="center"), None, None),
            (req_ids, Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (bva_ids, Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            (ast_ids, Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            ("", Alignment(horizontal="left", vertical="top"), None, None),              # Actual Result
            ("Chưa chạy", Alignment(horizontal="center", vertical="center"), None, None), # Execution status
            ("", Alignment(horizontal="left", vertical="top"), None, None),              # Notes
        ]

        for col_idx, (val, align, font_override, fill_override) in enumerate(row_values, 1):
            cell = ws_cases.cell(row=current_row, column=col_idx, value=val)
            cell.alignment = align
            cell.font = font_override or Font(name=FONT_NAME, size=9.5, color="0F172A")
            cell.fill = fill_override or row_fill
            cell.border = THIN_BORDER

        current_row += 1

    ws_cases.freeze_panes = "A6"
    if current_row > 6:
        ws_cases.auto_filter.ref = f"A5:Q{current_row - 1}"

    zero_widths = {
        1: 6,
        2: 16,
        3: 32,
        4: 14,
        5: 12,
        6: 18,
        7: 28,
        8: 26,
        9: 42,
        10: 34,
        11: 12,
        12: 20,
        13: 22,
        14: 22,
        15: 24,
        16: 16,
        17: 20,
    }
    _auto_fit_columns(ws_cases, len(headers_zero), zero_widths)

    # =========================================================================
    # SHEET 2: KHOẢNG TRỐNG QA (GAPS FOR QA)
    # =========================================================================
    if gaps_for_qa:
        ws_gaps = wb.create_sheet(title="Khoảng trống QA (Gaps)")
        ws_gaps.views.sheetView[0].showGridLines = True

        ws_gaps.merge_cells("A1:F1")
        top_gaps = ws_gaps["A1"]
        top_gaps.value = "DANH SÁCH KHOẢNG TRỐNG & XUNG ĐỘT SPEC CHUYỂN GIAO QA (RULE-06)"
        top_gaps.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
        top_gaps.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
        top_gaps.alignment = Alignment(horizontal="center", vertical="center")
        ws_gaps.row_dimensions[1].height = 36

        headers_gaps = [
            "STT",
            "Mã mục tiêu (Target ID)",
            "Loại khoảng trống (Gap Type)",
            "Lý do & Chi tiết kỹ thuật (Reason)",
            "Trạng thái phân loại (Status)",
            "Khuyến nghị hành động QA",
        ]

        ws_gaps.row_dimensions[3].height = 26
        for col_idx, header_text in enumerate(headers_gaps, 1):
            cell = ws_gaps.cell(row=3, column=col_idx, value=header_text)
            cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN_BORDER

        gap_row = 4
        for idx, gap in enumerate(gaps_for_qa, 1):
            is_even = (idx % 2 == 0)
            row_fill = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid") if is_even else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

            gap_type = gap.get("gap_type", "")
            action_rec = "Làm rõ với BA/PO về yêu cầu" if "SPEC_CONFLICT" in gap_type or "REQ" in gap_type else "Kiểm thử thủ công nhánh ẩn / biên API"

            gap_values = [
                (idx, Alignment(horizontal="center", vertical="center"), None, None),
                (gap.get("target_id", ""), Alignment(horizontal="center", vertical="center"), Font(name=FONT_NAME, size=9.5, bold=True), None),
                (gap_type, Alignment(horizontal="center", vertical="center"), Font(name=FONT_NAME, size=9.5, bold=True, color="B91C1C"), None),
                (gap.get("reason", ""), Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
                (gap.get("status", "NEEDS_CLARIFICATION"), Alignment(horizontal="center", vertical="center"), None, None),
                (action_rec, Alignment(horizontal="left", vertical="top", wrap_text=True), None, None),
            ]

            for col_idx, (val, align, font_override, fill_override) in enumerate(gap_values, 1):
                cell = ws_gaps.cell(row=gap_row, column=col_idx, value=val)
                cell.alignment = align
                cell.font = font_override or Font(name=FONT_NAME, size=9.5, color="0F172A")
                cell.fill = fill_override or row_fill
                cell.border = THIN_BORDER

            gap_row += 1

        ws_gaps.freeze_panes = "A4"
        if gap_row > 4:
            ws_gaps.auto_filter.ref = f"A3:F{gap_row - 1}"

        gap_widths = {1: 6, 2: 18, 3: 20, 4: 45, 5: 22, 6: 35}
        _auto_fit_columns(ws_gaps, len(headers_gaps), gap_widths)

    # Save
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
